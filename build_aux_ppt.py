import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_COLOR = RGBColor(248, 250, 252)        # Slate 50
    NAVY = RGBColor(15, 23, 42)               # Slate 900
    PRIMARY = RGBColor(37, 99, 235)           # Blue 600
    PRIMARY_LIGHT = RGBColor(239, 246, 255)   # Blue 50
    EMERALD = RGBColor(16, 185, 129)          # Emerald 500
    CARD_BG = RGBColor(255, 255, 255)         # White
    BORDER_COLOR = RGBColor(226, 232, 240)    # Slate 200
    TEXT_MUTED = RGBColor(100, 116, 139)      # Slate 500
    TEXT_DARK = RGBColor(30, 41, 59)          # Slate 800
    AMBER = RGBColor(217, 119, 6)             # Amber 600

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, subtitle_text=None):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.733), Inches(1.3))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_tag = tf.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.size = Pt(10.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = PRIMARY

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY
        p_title.space_before = Pt(3)

        if subtitle_text:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(12)
            p_sub.font.color.rgb = TEXT_MUTED
            p_sub.space_before = Pt(3)

    # =========================================================================
    # SLIDE 1: Problem Statement & Affected Users
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1)
    add_header(s1, "Executive Overview  •  Operational Context", 
               "AUX Break Management & Floor Adherence Breakdown",
               "The hidden cost of manual communication channels in modern contact center operations")

    # Top Hero Problem Statement Card
    p_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(11.733), Inches(1.9))
    p_card.fill.solid()
    p_card.fill.fore_color.rgb = CARD_BG
    p_card.line.color.rgb = RGBColor(203, 213, 225)
    p_card.line.width = Pt(1.5)

    p_tb = s1.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.133), Inches(1.6))
    ptf = p_tb.text_frame
    ptf.word_wrap = True
    ptf.margin_left = ptf.margin_top = ptf.margin_right = ptf.margin_bottom = 0

    p1_tag = ptf.paragraphs[0]
    p1_tag.text = "🚨  Core Problem Statement"
    p1_tag.font.size = Pt(14)
    p1_tag.font.bold = True
    p1_tag.font.color.rgb = RGBColor(220, 38, 38) # Red 600

    p1_body = ptf.add_paragraph()
    p1_body.text = (
        "Contact centers rely on manual AUX break requests and approvals through emails and group chats, "
        "resulting in communication delays, limited request visibility, inconsistent tracking, "
        "and difficulties maintaining agent adherence and workforce availability."
    )
    p1_body.font.size = Pt(13)
    p1_body.font.color.rgb = TEXT_DARK
    p1_body.space_before = Pt(6)

    # Bottom Grid: 4 Affected Users Cards
    user_lbl = s1.shapes.add_textbox(Inches(0.8), Inches(3.95), Inches(11.733), Inches(0.4))
    utf = user_lbl.text_frame
    utf.margin_left = utf.margin_top = utf.margin_right = utf.margin_bottom = 0
    p_u = utf.paragraphs[0]
    p_u.text = "👥  Affected Users Across the Operational Ecosystem"
    p_u.font.size = Pt(13)
    p_u.font.bold = True
    p_u.font.color.rgb = NAVY

    users = [
        ("Frontline Agents", "🎧", "Face uncertainty waiting on approvals, fragmented chat updates, and adherence penalties."),
        ("Team Leads & Supervisors", "👩‍💼", "Trapped in repetitive manual approvals and chat follow-ups instead of high-value coaching."),
        ("WFM Planners & Operations", "📊", "Struggle with invisible break schedules, unforecasted staffing deficits, and adherence leaks."),
        ("End Customers", "🤝", "Endure longer queue wait times and delayed support caused by sudden staffing gaps.")
    ]

    card_w = Inches(2.78)
    card_h = Inches(2.3)
    spacing = Inches(0.2)
    start_x = Inches(0.8)

    for i, (role, icon, detail) in enumerate(users):
        x = start_x + i * (card_w + spacing)
        u_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(4.45), card_w, card_h)
        u_card.fill.solid()
        u_card.fill.fore_color.rgb = CARD_BG
        u_card.line.color.rgb = BORDER_COLOR
        u_card.line.width = Pt(1.5)

        tb = s1.shapes.add_textbox(x + Inches(0.18), Inches(4.6), card_w - Inches(0.36), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_icon = tf.paragraphs[0]
        p_icon.text = f"{icon}  {role}"
        p_icon.font.size = Pt(13)
        p_icon.font.bold = True
        p_icon.font.color.rgb = PRIMARY

        p_det = tf.add_paragraph()
        p_det.text = detail
        p_det.font.size = Pt(10.5)
        p_det.font.color.rgb = TEXT_MUTED
        p_det.space_before = Pt(8)

    s1.notes_slide.notes_text_frame.text = (
        "SLIDE 1 TALK TRACK (0:00 - 0:35):\n"
        "Contact centers today still rely on manual AUX break requests and approvals through emails and group chats. "
        "This legacy approach causes severe communication delays, zero real-time request visibility, inconsistent tracking, "
        "and major breakdowns in agent adherence and floor availability. "
        "Four critical stakeholder groups are directly impacted: Frontline Agents, Team Leads, WFM Operations, and ultimately End Customers."
    )

    # =========================================================================
    # SLIDE 2: Pain Points – AUX Break Management
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2)
    add_header(s2, "Operational Bottlenecks  •  Root Causes",
               "Pain Points – AUX Break Management",
               "How manual workflows degrade efficiency across agents, supervisors, WFM, and customers")

    pain_points = [
        ("Agents", "🎧", "Delayed break approvals, unclear request status, and difficulty communicating break requests.",
         "Impact: Idle waiting off-queue, schedule frustration, and compliance confusion."),
        ("Team Leaders (TLs)", "👩‍💼", "High volume of manual requests, repetitive follow-ups, and difficulty tracking approvals.",
         "Impact: 1.5–2 hours wasted daily on routine approvals instead of frontline coaching."),
        ("WFM Management", "📊", "Limited visibility into agent AUX status, break schedules, and real-time staffing availability.",
         "Impact: Intraday staffing variance, schedule leakage, and reactive overtime costs."),
        ("End Customers", "🤝", "Longer wait times and delayed support due to reduced agent availability and staffing gaps.",
         "Impact: Increased queue abandonment, rushed interactions, and lower CSAT scores.")
    ]

    grid_w = Inches(5.72)
    grid_h = Inches(2.25)
    row_gap = Inches(0.25)
    col_gap = Inches(0.29)

    for i, (persona, icon, pain, sub_impact) in enumerate(pain_points):
        col = i % 2
        row = i // 2
        x = Inches(0.8) + col * (grid_w + col_gap)
        y = Inches(1.95) + row * (grid_h + row_gap)

        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, grid_w, grid_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(239, 68, 68) if col == 0 else RGBColor(217, 119, 6)
        card.line.width = Pt(1.5)

        tb = s2.shapes.add_textbox(x + Inches(0.25), y + Inches(0.18), grid_w - Inches(0.5), grid_h - Inches(0.36))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_h = tf.paragraphs[0]
        p_h.text = f"{icon}  {persona}"
        p_h.font.size = Pt(14)
        p_h.font.bold = True
        p_h.font.color.rgb = NAVY

        p_pain = tf.add_paragraph()
        p_pain.text = f"• {pain}"
        p_pain.font.size = Pt(11.5)
        p_pain.font.color.rgb = TEXT_DARK
        p_pain.space_before = Pt(6)

        p_imp = tf.add_paragraph()
        p_imp.text = sub_impact
        p_imp.font.size = Pt(10)
        p_imp.font.color.rgb = TEXT_MUTED
        p_imp.space_before = Pt(6)

    s2.notes_slide.notes_text_frame.text = (
        "SLIDE 2 TALK TRACK (0:35 - 1:15):\n"
        "Here are the specific pain points we uncover on the floor:\n"
        "• Agents experience delayed break approvals, unclear request status, and friction communicating requests.\n"
        "• Team Leaders are bogged down by high volumes of manual requests, constant follow-ups, and cumbersome tracking.\n"
        "• WFM Management lacks real-time visibility into agent AUX statuses, break schedules, and true staffing availability.\n"
        "• End Customers suffer the downstream consequences: longer queue wait times and delayed support caused by sudden staffing gaps."
    )

    # =========================================================================
    # SLIDE 3: Business Benefits & Measurable Value
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3)
    add_header(s3, "Business Benefits  •  Measurable ROI",
               "Tangible Business Benefits of Automated AUX Management",
               "Recovering capacity, restoring supervisor coaching time, and protecting customer SLA commitments")

    benefits = [
        ("Eliminating Break Delays & Confusion", "⏱️",
         "Recovers 15–20 minutes of productive shift time per agent daily",
         "Instant self-service evaluation eliminates off-queue waiting time and improves agent satisfaction."),
        
        ("Eliminating Manual Request Overload", "👩‍💼",
         "Automating routine approvals reclaims up to two hours of supervisory capacity per shift",
         "Frees Team Leaders from constant chat interruptions to focus on coaching and escalation resolution."),
        
        ("Eliminating Visibility Gaps & Staffing Variance", "📊",
         "Eliminates unplanned overtime costs through 100% real-time floor visibility",
         "Synchronized CTI telemetry and wallboards ensure strict schedule adherence and predictable staffing."),
        
        ("Eliminating Long Wait Times & Staffing Deficits", "🤝",
         "Reduces queue wait times (ASA) by 15–25% and curbs call abandonment during critical call volume peaks",
         "Intelligent break throttling prevents simultaneous break stacking during peak caller arrival windows.")
    ]

    for i, (title, icon, benefit_stat, impact_detail) in enumerate(benefits):
        col = i % 2
        row = i // 2
        x = Inches(0.8) + col * (grid_w + col_gap)
        y = Inches(1.95) + row * (grid_h + row_gap)

        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, grid_w, grid_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = EMERALD if row == 0 else PRIMARY
        card.line.width = Pt(1.8)

        tb = s3.shapes.add_textbox(x + Inches(0.25), y + Inches(0.18), grid_w - Inches(0.5), grid_h - Inches(0.36))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_h = tf.paragraphs[0]
        p_h.text = f"{icon}  {title}"
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = NAVY

        p_stat = tf.add_paragraph()
        p_stat.text = f"★  {benefit_stat}"
        p_stat.font.size = Pt(12)
        p_stat.font.bold = True
        p_stat.font.color.rgb = EMERALD if row == 0 else PRIMARY
        p_stat.space_before = Pt(6)

        p_det = tf.add_paragraph()
        p_det.text = impact_detail
        p_det.font.size = Pt(10)
        p_det.font.color.rgb = TEXT_MUTED
        p_det.space_before = Pt(5)

    s3.notes_slide.notes_text_frame.text = (
        "SLIDE 3 TALK TRACK (1:15 - 1:55):\n"
        "By replacing manual processes with automated AUX break governance, the business realizes 4 core benefits:\n"
        "1. Eliminating Break Delays & Confusion recovers 15 to 20 minutes of productive shift time per agent daily.\n"
        "2. Eliminating Manual Request Overload reclaims up to two hours of supervisory capacity per shift.\n"
        "3. Eliminating Visibility Gaps & Staffing Variance eliminates unplanned overtime costs through 100% real-time floor visibility.\n"
        "4. Eliminating Long Wait Times & Staffing Deficits reduces queue wait times by 15 to 25% and curbs call abandonment during critical call volume peaks.\n"
        "This turns chaotic break management into measurable, automated operational ROI."
    )

    out_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "AUX_Break_Management_Pitch.pptx")
    prs.save(out_file)
    print("Successfully generated:", out_file)

if __name__ == "__main__":
    build_presentation()
