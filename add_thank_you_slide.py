import pptx
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR

# Load clean backup containing original 3 slides
prs = pptx.Presentation("C:/Users/maddu/Downloads/AUX_Break_Management_backup.pptx")

# 1. Update Slide 1 with ZEL App Name and Tagline
slide1 = prs.slides[0]
for sh in slide1.shapes:
    if sh.name == "Text 0":
        p1 = sh.text_frame.paragraphs[0]
        p1.text = "Problem Statement"
        p1.font.size = Pt(32)
        p1.font.bold = True
        
        p2 = sh.text_frame.add_paragraph()
        p2.space_before = Pt(4)
        
        run_name = p2.add_run()
        run_name.text = "ZEL"
        run_name.font.bold = True
        run_name.font.size = Pt(14)
        
        run_sep = p2.add_run()
        run_sep.text = " \u2014 "
        run_sep.font.size = Pt(14)
        
        run_tag = p2.add_run()
        run_tag.text = '"Zero wait to pause, Zero lag to serve"'
        run_tag.font.italic = True
        run_tag.font.size = Pt(13)

# 2. Add Slide 4: Thank You Slide (Matching Exact Visual Design of Slides 1-3)
layout = prs.slide_layouts[1] # CONTENT layout
slide4 = prs.slides.add_slide(layout)

# Top Right Corner Motifs (Exact match to Slides 1-3)
m1 = slide4.shapes.add_shape(MSO_SHAPE.OVAL, 8001000, 365760, 365760, 365760)
m1.fill.solid()
m1.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
m1.line.fill.background()

m2 = slide4.shapes.add_shape(MSO_SHAPE.OVAL, 8211312, 365760, 365760, 365760)
m2.fill.solid()
m2.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
m2.line.fill.background()

# Eyebrow Tag
eb = slide4.shapes.add_textbox(548640, 320040, 6400800, 274320)
eb_p = eb.text_frame.paragraphs[0]
eb_p.text = "04 \u00b7 CONCLUSION"
eb_r = eb_p.runs[0]
eb_r.font.size = Pt(12)
eb_r.font.bold = True
eb_r.font.color.theme_color = MSO_THEME_COLOR.TEXT_1

# Main Title & Subtitle
title_shape = None
for sh in slide4.shapes:
    if sh.name == "Text 0":
        title_shape = sh
        break

if not title_shape:
    title_shape = slide4.shapes.add_textbox(548640, 640080, 8046720, 731520)

tf = title_shape.text_frame
p1 = tf.paragraphs[0]
p1.text = "Thank You"
p1.font.size = Pt(36)
p1.font.bold = True
p1.font.color.theme_color = MSO_THEME_COLOR.TEXT_1

p2 = tf.add_paragraph()
p2.space_before = Pt(4)
r_name = p2.add_run()
r_name.text = "ZEL"
r_name.font.bold = True
r_name.font.size = Pt(14)
r_name.font.color.theme_color = MSO_THEME_COLOR.TEXT_1

r_sep = p2.add_run()
r_sep.text = " \u2014 "
r_sep.font.size = Pt(14)
r_sep.font.color.theme_color = MSO_THEME_COLOR.TEXT_1

r_tag = p2.add_run()
r_tag.text = '"Zero wait to pause, Zero lag to serve"'
r_tag.font.italic = True
r_tag.font.size = Pt(13)
r_tag.font.color.theme_color = MSO_THEME_COLOR.TEXT_1

# Main Container Card
card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 548640, 1463040, 8046720, 3200400)
card.fill.solid()
card.fill.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
card.line.color.theme_color = MSO_THEME_COLOR.ACCENT_1
card.line.width = Pt(1.5)

# Left Column: Autonomous Capabilities
left_tb = slide4.shapes.add_textbox(822960, 1691640, 3840480, 2743200)
ltf = left_tb.text_frame
ltf.word_wrap = True

lp0 = ltf.paragraphs[0]
lp0.text = "Autonomous Floor Operations"
lp0.font.size = Pt(16)
lp0.font.bold = True
lp0.font.color.theme_color = MSO_THEME_COLOR.TEXT_1

lp1 = ltf.add_paragraph()
lp1.text = "ZEL transforms contact center break management from manual chat delays into instant, algorithmic workforce flow."
lp1.font.size = Pt(11)
lp1.space_before = Pt(8)
lp1.font.color.theme_color = MSO_THEME_COLOR.TEXT_2

pillars = [
    ("\u2022 Instant AUX Approvals: ", "Sub-second Erlang-C staffing evaluation."),
    ("\u2022 Autonomous ACW: ", "AI summary generation & direct CRM sync."),
    ("\u2022 Dynamic Floor Flex: ", "1-click intraday queue rebalancing.")
]
for p_title, p_desc in pillars:
    p = ltf.add_paragraph()
    p.space_before = Pt(6)
    r1 = p.add_run()
    r1.text = p_title
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r1.font.color.theme_color = MSO_THEME_COLOR.TEXT_1
    r2 = p.add_run()
    r2.text = p_desc
    r2.font.size = Pt(10.5)
    r2.font.color.theme_color = MSO_THEME_COLOR.TEXT_2

# Right Column: Business Outcomes Recap & Discussion
right_tb = slide4.shapes.add_textbox(4846320, 1691640, 3474720, 2743200)
rtf = right_tb.text_frame
rtf.word_wrap = True

rp0 = rtf.paragraphs[0]
rp0.text = "Business Outcomes Delivered"
rp0.font.size = Pt(16)
rp0.font.bold = True
rp0.font.color.theme_color = MSO_THEME_COLOR.TEXT_1

outcomes = [
    ("\u2022 15\u201320 min: ", "Productive shift time recovered per agent daily"),
    ("\u2022 2 hrs: ", "Supervisor capacity reclaimed per shift"),
    ("\u2022 100%: ", "Real-time floor visibility & zero overtime costs"),
    ("\u2022 15\u201325%: ", "Lower queue wait times during volume surges")
]

for stat, desc in outcomes:
    p = rtf.add_paragraph()
    p.space_before = Pt(5)
    r1 = p.add_run()
    r1.text = stat
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r1.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(10.5)
    r2.font.color.theme_color = MSO_THEME_COLOR.TEXT_2

p_qa = rtf.add_paragraph()
p_qa.space_before = Pt(14)
r_qa = p_qa.add_run()
r_qa.text = "Questions & Discussion"
r_qa.font.bold = True
r_qa.font.size = Pt(14)
r_qa.font.color.theme_color = MSO_THEME_COLOR.ACCENT_2

# Speaker Notes for Slide 4
slide4.notes_slide.notes_text_frame.text = (
    "CLOSING TALK TRACK:\n"
    "Thank you for your time. "
    "With ZEL — 'Zero wait to pause, Zero lag to serve' — contact centers eliminate the daily friction of manual break approvals, "
    "empower agents with autonomous ACW, and protect SLA commitments with real-time floor balancing. "
    "I'd be glad to take any questions and discuss next steps."
)

# Save files
out_paths = [
    "C:/Users/maddu/Downloads/AUX_Break_Management.pptx",
    "C:/Users/maddu/.gemini/antigravity/scratch/aux-chatbot-/AUX_Break_Management.pptx"
]

for p in out_paths:
    prs.save(p)
    print("Successfully saved:", p)
